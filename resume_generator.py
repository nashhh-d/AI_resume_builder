import io
import base64

from docx import Document
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet


# -----------------------------
# DOCX RESUME GENERATOR
# -----------------------------

def build_docx(data, summary, suggested_skills):

    doc = Document()

    # Name
    doc.add_heading(data.get("name", ""), 0)

    # Contact
    contact = f'{data.get("email","")} | {data.get("phone","")} | {data.get("linkedin","")}'
    doc.add_paragraph(contact)

    # Summary
    doc.add_heading("Summary", level=1)
    doc.add_paragraph(summary)

    # Skills
    all_skills = data.get("skills", []) + suggested_skills
    if all_skills:
        doc.add_heading("Skills", level=1)
        doc.add_paragraph(", ".join(all_skills))

    # Experience
    experience = data.get("experience", [])
    if experience:
        doc.add_heading("Experience", level=1)

        for exp in experience:
            role_line = f'{exp.get("role","")} - {exp.get("company","")} ({exp.get("duration","")})'
            doc.add_paragraph(role_line, style="List Bullet")

            description = exp.get("description", "")
            if description:
                doc.add_paragraph(description)

    # Education
    education = data.get("education", [])
    if education:
        doc.add_heading("Education", level=1)

        for edu in education:
            line = f'{edu.get("degree","")} - {edu.get("institution","")} ({edu.get("year","")})'
            doc.add_paragraph(line, style="List Bullet")

    # Projects
    projects = data.get("projects", [])
    if projects:
        doc.add_heading("Projects", level=1)

        for proj in projects:
            line = f'{proj.get("title","")}: {proj.get("description","")}'
            doc.add_paragraph(line, style="List Bullet")

    # Certifications
    certs = data.get("certifications", [])
    if certs:
        doc.add_heading("Certifications", level=1)

        for cert in certs:
            doc.add_paragraph(cert, style="List Bullet")

    # Save DOCX
    buffer = io.BytesIO()
    doc.save(buffer)

    docx_base64 = base64.b64encode(buffer.getvalue()).decode()

    return docx_base64


# -----------------------------
# PDF RESUME GENERATOR
# -----------------------------

def build_pdf(data, summary, suggested_skills):

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer)

    styles = getSampleStyleSheet()
    elements = []

    # Name
    elements.append(Paragraph(data.get("name",""), styles["Title"]))

    # Contact
    contact = f'{data.get("email","")} | {data.get("phone","")} | {data.get("linkedin","")}'
    elements.append(Paragraph(contact, styles["Normal"]))
    elements.append(Spacer(1, 10))

    # Summary
    elements.append(Paragraph("<b>Summary</b>", styles["Heading2"]))
    elements.append(Paragraph(summary, styles["Normal"]))
    elements.append(Spacer(1, 10))

    # Skills
    all_skills = data.get("skills", []) + suggested_skills
    if all_skills:
        elements.append(Paragraph("<b>Skills</b>", styles["Heading2"]))
        elements.append(Paragraph(", ".join(all_skills), styles["Normal"]))
        elements.append(Spacer(1, 10))

    # Experience
    experience = data.get("experience", [])
    if experience:
        elements.append(Paragraph("<b>Experience</b>", styles["Heading2"]))

        for exp in experience:
            role_line = f'{exp.get("role","")} - {exp.get("company","")} ({exp.get("duration","")})'
            elements.append(Paragraph(role_line, styles["Normal"]))

            description = exp.get("description", "")
            if description:
                elements.append(Paragraph(description, styles["Normal"]))

        elements.append(Spacer(1, 10))

    # Education
    education = data.get("education", [])
    if education:
        elements.append(Paragraph("<b>Education</b>", styles["Heading2"]))

        for edu in education:
            line = f'{edu.get("degree","")} - {edu.get("institution","")} ({edu.get("year","")})'
            elements.append(Paragraph(line, styles["Normal"]))

        elements.append(Spacer(1, 10))

    # Projects
    projects = data.get("projects", [])
    if projects:
        elements.append(Paragraph("<b>Projects</b>", styles["Heading2"]))

        for proj in projects:
            line = f'{proj.get("title","")}: {proj.get("description","")}'
            elements.append(Paragraph(line, styles["Normal"]))

        elements.append(Spacer(1, 10))

    # Certifications
    certs = data.get("certifications", [])
    if certs:
        elements.append(Paragraph("<b>Certifications</b>", styles["Heading2"]))

        for cert in certs:
            elements.append(Paragraph(cert, styles["Normal"]))

    doc.build(elements)

    pdf_base64 = base64.b64encode(buffer.getvalue()).decode()

    return pdf_base64
